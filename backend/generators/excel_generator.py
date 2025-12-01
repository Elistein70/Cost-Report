"""Excel report generator - creates final DOH Cost Report"""
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from pathlib import Path
from config import settings
from models import Report, FlaggedItem
from processors.payroll_processor import PayrollProcessor
from processors.trial_balance_processor import TrialBalanceProcessor
from processors.visit_processor import VisitProcessor


class ExcelGenerator:
    """Generate final DOH Cost Report Excel file"""

    def __init__(self, report: Report, db_session):
        self.report = report
        self.db = db_session
        self.agency = report.agency
        self.template_path = settings.TEMPLATE_DIR / "Core_Schedules_Template_2024.xlsx"
        self.output_path = None

        # Color definitions (matching template)
        self.colors = {
            'LightBlue': 'CCFFFF',
            'Red': 'FFCCCC',
            'Orange': 'FFE5CC',
            'Purple': 'E6CCFF',
        }

    def generate(self) -> str:
        """Generate complete Excel report"""
        # Step 1: Copy template
        output_filename = f"{self.agency.name}_DOH_Cost_Report_{self.agency.year}_FINAL.xlsx"
        self.output_path = settings.OUTPUT_DIR / output_filename

        # Check if template exists
        if not self.template_path.exists():
            # Create a basic template structure
            self._create_basic_template()
        else:
            # Copy existing template
            import shutil
            shutil.copy(self.template_path, self.output_path)

        # Step 2: Fill in data
        self._fill_payroll_data()
        self._fill_trial_balance_data()
        self._fill_visit_data()

        # Step 3: Apply user corrections from flagged items
        self._apply_user_corrections()

        # Step 4: Format cells (colors, etc.)
        self._apply_formatting()

        # Step 5: Refresh formulas/pivots (in production, use win32com or xlwings)
        # For now, just save
        print(f"Generated report: {self.output_path}")

        return str(self.output_path)

    def _create_basic_template(self):
        """Create basic template structure if template file doesn't exist"""
        with pd.ExcelWriter(self.output_path, engine='openpyxl') as writer:
            # Create basic sheets
            pd.DataFrame().to_excel(writer, sheet_name='Cover', index=False)
            pd.DataFrame().to_excel(writer, sheet_name='Payroll Reports Input', index=False)
            pd.DataFrame().to_excel(writer, sheet_name='TB', index=False)
            pd.DataFrame().to_excel(writer, sheet_name='TB Tagging', index=False)
            pd.DataFrame().to_excel(writer, sheet_name='Detail Data 5 Tagging', index=False)
            pd.DataFrame().to_excel(writer, sheet_name='Schedule 11', index=False)
            pd.DataFrame().to_excel(writer, sheet_name='Reconciliation', index=False)

    def _fill_payroll_data(self):
        """Fill payroll data into Excel"""
        if not self.report.payroll_file:
            return

        # Process payroll file
        processor = PayrollProcessor(self.report.payroll_file, self.db)
        processor.process()

        # Load workbook
        wb = load_workbook(self.output_path)

        # Write to Payroll Reports Input sheet
        if 'Payroll Reports Input' in wb.sheetnames:
            ws = wb['Payroll Reports Input']
            # Clear existing data (keep headers)
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                for cell in row:
                    cell.value = None

            # Write new data
            if processor.tagged_data is not None:
                for r_idx, row in processor.tagged_data.iterrows():
                    excel_row = r_idx + 2  # +2 for header row
                    ws.cell(row=excel_row, column=1, value=row.get('Department_L1'))
                    ws.cell(row=excel_row, column=2, value=row.get('Department_L2'))
                    ws.cell(row=excel_row, column=3, value=row.get('Department_L3'))
                    ws.cell(row=excel_row, column=4, value=row.get('Department Name'))
                    ws.cell(row=excel_row, column=5, value=row.get('Paycode_Allocation'))
                    ws.cell(row=excel_row, column=6, value=row.get('Earning Code'))
                    ws.cell(row=excel_row, column=7, value=row.get('Hours'))
                    ws.cell(row=excel_row, column=8, value=row.get('Amount'))
                    ws.cell(row=excel_row, column=11, value=row.get('Tax_Allocation'))
                    ws.cell(row=excel_row, column=12, value=row.get('ER Tax Code'))
                    ws.cell(row=excel_row, column=13, value=row.get('ER Tax Amount'))

        wb.save(self.output_path)

    def _fill_trial_balance_data(self):
        """Fill trial balance data into Excel"""
        if not self.report.trial_balance_file:
            return

        processor = TrialBalanceProcessor(self.report.trial_balance_file, self.db)
        processor.process()

        wb = load_workbook(self.output_path)

        # Write raw TB
        if 'TB' in wb.sheetnames:
            ws = wb['TB']
            for r_idx, row in processor.data.iterrows():
                excel_row = r_idx + 2
                ws.cell(row=excel_row, column=1, value=row.get('Account'))
                ws.cell(row=excel_row, column=2, value=row.get('Description'))
                ws.cell(row=excel_row, column=3, value=row.get('Balance'))

        # Write tagged TB
        if 'TB Tagging' in wb.sheetnames:
            ws = wb['TB Tagging']
            for r_idx, row in processor.tagged_data.iterrows():
                excel_row = r_idx + 2
                ws.cell(row=excel_row, column=1, value=row.get('Account'))
                ws.cell(row=excel_row, column=2, value=row.get('Description'))
                ws.cell(row=excel_row, column=3, value=row.get('Balance'))
                ws.cell(row=excel_row, column=4, value=row.get('Program'))
                ws.cell(row=excel_row, column=5, value=row.get('Category'))
                ws.cell(row=excel_row, column=6, value=row.get('Router'))

        wb.save(self.output_path)

    def _fill_visit_data(self):
        """Fill visit data into Excel"""
        if not self.report.visit_file:
            return

        processor = VisitProcessor(self.report.visit_file, self.db)
        processor.process()

        wb = load_workbook(self.output_path)

        if 'Detail Data 5 Tagging' in wb.sheetnames:
            ws = wb['Detail Data 5 Tagging']
            for r_idx, row in processor.tagged_data.iterrows():
                excel_row = r_idx + 2
                ws.cell(row=excel_row, column=1, value=row.get('Billing Code'))
                ws.cell(row=excel_row, column=2, value=row.get('Service Description'))
                ws.cell(row=excel_row, column=3, value=row.get('Program_Type'))
                ws.cell(row=excel_row, column=4, value=row.get('Visit Count'))
                ws.cell(row=excel_row, column=5, value=row.get('Billed Hours'))
                ws.cell(row=excel_row, column=6, value=row.get('County'))
                ws.cell(row=excel_row, column=7, value=row.get('Payer'))

        wb.save(self.output_path)

    def _apply_user_corrections(self):
        """Apply user corrections from reviewed flagged items"""
        flagged_items = self.db.query(FlaggedItem).filter(
            FlaggedItem.report_id == self.report.id,
            FlaggedItem.reviewed == True
        ).all()

        wb = load_workbook(self.output_path)

        for item in flagged_items:
            if item.user_corrected_tag and item.sheet_name in wb.sheetnames:
                ws = wb[item.sheet_name]
                # Apply correction at row_number
                # (column depends on item_type - simplified here)
                if item.row_number:
                    ws.cell(row=item.row_number, column=5, value=item.user_corrected_tag)

        wb.save(self.output_path)

    def _apply_formatting(self):
        """Apply color coding and formatting"""
        wb = load_workbook(self.output_path)

        # Color code TB Tagging sheet (payroll allocation lines)
        if 'TB Tagging' in wb.sheetnames:
            ws = wb['TB Tagging']
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                category = row[4].value  # Column E (Category)
                if category in ['Salaries', 'Payroll Taxes', 'Fringe Benefits', 'Workers Compensation']:
                    # Apply color to entire row
                    color_name = {
                        'Salaries': 'LightBlue',
                        'Payroll Taxes': 'Red',
                        'Fringe Benefits': 'Orange',
                        'Workers Compensation': 'Purple',
                    }.get(category)

                    if color_name and color_name in self.colors:
                        fill = PatternFill(start_color=self.colors[color_name],
                                         end_color=self.colors[color_name],
                                         fill_type='solid')
                        for cell in row:
                            cell.fill = fill

        wb.save(self.output_path)
