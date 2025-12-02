"""Payroll file processor - Updated for Anchor and Nemo template structure"""
import pandas as pd
import pdfplumber
from pathlib import Path
from typing import List, Dict, Any
from rapidfuzz import fuzz
try:
    from config import settings
except ImportError:
    from backend.config_simple import settings


class PayrollProcessor:
    """Process payroll register files (PDF or CSV) following SOP rules"""

    def __init__(self, file_path: str, db_session):
        self.file_path = Path(file_path)
        self.db = db_session
        self.data = None
        self.tagged_data = None
        self.flagged_items = []
        self.questions = []  # Questions for user clarification

        # Master payroll allocation guide
        # Based on actual template analysis - Column E mappings
        self.paycode_tags = {
            "Regular": "Base Wages for Hours Worked",
            "Regular Pay": "Base Wages for Hours Worked",
            "Reg Pay": "Base Wages for Hours Worked",
            "Straight Time": "Base Wages for Hours Worked",

            "Overtime": "Overtime Wages",
            "OT": "Overtime Wages",
            "OT Pay": "Overtime Wages",

            "Bonus": "Premium Pay",
            "Premium": "Premium Pay",

            "Meal": "Other Wages",
            "Misc": "Other Wages",
            "Misc Pay": "Other Wages",

            "Holiday": "Holiday Pay",
            "Holiday Pay": "Holiday Pay",

            "Sick": "Sick Time",
            "Sick Pay": "Sick Time",

            "Vacation": "Vacation Time",
            "PTO": "PTO",
        }

        # Tax code mappings - Column L
        self.tax_tags = {
            "SS-R": "FICA Taxes",
            "FICA": "FICA Taxes",
            "MED-R": "FICA Taxes",
            "Medicare": "FICA Taxes",

            "FUTA": "Disability/Unemployment/Workers Compensation Taxes",
            "NYSUI": "Disability/Unemployment/Workers Compensation Taxes",
            "SUI": "Disability/Unemployment/Workers Compensation Taxes",
            "NY-MTA1": "Disability/Unemployment/Workers Compensation Taxes",
            "NYCLA": "Disability/Unemployment/Workers Compensation Taxes",

            "WC": "Workers Compensation Taxes",
            "Workers Comp": "Workers Compensation Taxes",
        }

    def process(self) -> List[Dict[str, Any]]:
        """Main processing method following SOP"""
        # Step 1: Extract data from file
        if self.file_path.suffix.lower() == '.pdf':
            self.data = self._extract_from_pdf()
        else:
            self.data = self._extract_from_csv()

        if self.data is None or self.data.empty:
            self.questions.append("Payroll file appears to be empty or unreadable. Please verify the file format.")
            return self.flagged_items

        # Step 2: Auto-tag the data
        self.tagged_data = self._auto_tag()

        # Step 3: Return flagged items for review
        return self.flagged_items

    def _extract_from_pdf(self) -> pd.DataFrame:
        """Extract payroll data from PDF - following SOP parsing rules"""
        rows = []
        current_department = None

        try:
            with pdfplumber.open(self.file_path) as pdf:
                for page_num, page in enumerate(pdf.pages, 1):
                    text = page.extract_text()
                    if not text:
                        continue

                    # Parse line by line
                    for line in text.split('\n'):
                        line = line.strip()
                        if not line:
                            continue

                        # Detect department headers
                        if any(keyword in line.upper() for keyword in ['DEPARTMENT', 'DEPT']):
                            current_department = line
                            continue

                        # Try to parse payroll data lines
                        # Format varies by provider (Viventium, Empeon, ADP, etc.)
                        parts = line.split()

                        if len(parts) >= 3:
                            try:
                                # Attempt to extract: Code, Hours, Amount
                                # This is a simplified parser - adjust based on your PDF format
                                code = parts[0]
                                hours = None
                                amount = None

                                # Find numeric values
                                for i, part in enumerate(parts):
                                    part_clean = part.replace('$', '').replace(',', '')
                                    try:
                                        val = float(part_clean)
                                        if hours is None and val < 100000:  # Likely hours
                                            hours = val
                                        elif amount is None:  # Likely amount
                                            amount = val
                                    except ValueError:
                                        continue

                                if code and (hours is not None or amount is not None):
                                    rows.append({
                                        'Department': current_department or 'Unknown',
                                        'Code': code,
                                        'Hours': hours or 0,
                                        'Amount': amount or 0,
                                        'Tax_Code': None,
                                        'Tax_Amount': 0
                                    })
                            except (ValueError, IndexError):
                                continue

        except Exception as e:
            self.questions.append(f"Error reading PDF: {str(e)}. Please provide payroll data in Excel/CSV format instead.")
            return pd.DataFrame()

        if not rows:
            self.questions.append("Could not extract payroll data from PDF. The format may not be recognized. Please provide an Excel or CSV export.")

        return pd.DataFrame(rows)

    def _extract_from_csv(self) -> pd.DataFrame:
        """Extract payroll data from CSV/Excel"""
        try:
            if self.file_path.suffix.lower() in ['.xlsx', '.xls', '.xlsm']:
                df = pd.read_excel(self.file_path, engine='openpyxl')
            else:
                df = pd.read_csv(self.file_path)

            # Standardize column names (case-insensitive matching)
            column_mapping = {}
            for col in df.columns:
                col_lower = str(col).strip().lower()

                if any(x in col_lower for x in ['dept', 'department']):
                    column_mapping[col] = 'Department'
                elif any(x in col_lower for x in ['code', 'earning', 'paycode']):
                    column_mapping[col] = 'Code'
                elif any(x in col_lower for x in ['hrs', 'hours']):
                    column_mapping[col] = 'Hours'
                elif any(x in col_lower for x in ['amount', 'gross', 'wages']):
                    column_mapping[col] = 'Amount'
                elif any(x in col_lower for x in ['tax code', 'tax_code', 'er tax']):
                    column_mapping[col] = 'Tax_Code'
                elif any(x in col_lower for x in ['tax amount', 'tax_amount', 'er amount']):
                    column_mapping[col] = 'Tax_Amount'

            df = df.rename(columns=column_mapping)

            # Ensure required columns exist
            for col in ['Department', 'Code', 'Hours', 'Amount']:
                if col not in df.columns:
                    df[col] = '' if col == 'Department' or col == 'Code' else 0

            if 'Tax_Code' not in df.columns:
                df['Tax_Code'] = None
            if 'Tax_Amount' not in df.columns:
                df['Tax_Amount'] = 0

            return df

        except Exception as e:
            self.questions.append(f"Error reading file: {str(e)}. Please check the file format.")
            return pd.DataFrame()

    def _auto_tag(self) -> pd.DataFrame:
        """Auto-tag payroll data following SOP allocation rules"""
        df = self.data.copy()

        # Add tagging columns (matching template structure)
        df['Tag'] = ''  # Column E - Paycode allocation
        df['Tax_Category'] = ''  # Column K - Tax category
        df['Confidence_Paycode'] = 0.0
        df['Confidence_Tax'] = 0.0

        for idx, row in df.iterrows():
            # Tag earning code (Column E/F -> Column E tag)
            paycode_tag, paycode_confidence = self._tag_paycode(row.get('Code', ''))
            df.at[idx, 'Tag'] = paycode_tag
            df.at[idx, 'Confidence_Paycode'] = paycode_confidence

            # Flag if confidence is low
            if paycode_confidence < settings.CONFIDENCE_THRESHOLD and paycode_tag != 'Unknown':
                self.flagged_items.append({
                    'original_value': row.get('Code', ''),
                    'suggested_tag': paycode_tag,
                    'confidence': paycode_confidence,
                    'row_number': idx + 6  # Template data starts at row 6
                })
            elif paycode_tag == 'Unknown':
                self.questions.append(f"Unrecognized payroll code: '{row.get('Code', '')}'. What should this be tagged as?")

            # Tag tax code
            if pd.notna(row.get('Tax_Code')) and row.get('Tax_Code'):
                tax_tag, tax_confidence = self._tag_tax_code(row['Tax_Code'])
                df.at[idx, 'Tax_Category'] = tax_tag
                df.at[idx, 'Confidence_Tax'] = tax_confidence

                if tax_confidence < settings.CONFIDENCE_THRESHOLD and tax_tag != 'Unknown':
                    self.flagged_items.append({
                        'original_value': row['Tax_Code'],
                        'suggested_tag': tax_tag,
                        'confidence': tax_confidence,
                        'row_number': idx + 6
                    })
                elif tax_tag == 'Unknown':
                    self.questions.append(f"Unrecognized tax code: '{row['Tax_Code']}'. What should this be tagged as?")

        return df

    def _tag_paycode(self, code: str) -> tuple:
        """Tag paycode with fuzzy matching"""
        if not code or pd.isna(code):
            return 'Unknown', 0.5

        code_clean = str(code).strip()

        # Exact match
        if code_clean in self.paycode_tags:
            return self.paycode_tags[code_clean], 1.0

        # Fuzzy match
        best_match = None
        best_score = 0
        for key, value in self.paycode_tags.items():
            score = fuzz.ratio(code_clean.lower(), key.lower()) / 100
            if score > best_score:
                best_score = score
                best_match = value

        if best_match and best_score > 0.75:
            return best_match, best_score

        return 'Unknown', 0.5

    def _tag_tax_code(self, tax_code: str) -> tuple:
        """Tag tax code with fuzzy matching"""
        if not tax_code or pd.isna(tax_code):
            return 'Unknown', 0.5

        tax_clean = str(tax_code).strip()

        # Exact match
        if tax_clean in self.tax_tags:
            return self.tax_tags[tax_clean], 1.0

        # Fuzzy match
        best_match = None
        best_score = 0
        for key, value in self.tax_tags.items():
            score = fuzz.ratio(tax_clean.lower(), key.lower()) / 100
            if score > best_score:
                best_score = score
                best_match = value

        if best_match and best_score > 0.75:
            return best_match, best_score

        return 'Unknown', 0.5

    def save_to_template(self, template_path: str, output_path: str):
        """
        Save tagged data to the actual template file
        Writes to columns E, F, G, I, L, M starting at row 6
        """
        from openpyxl import load_workbook

        if self.tagged_data is None:
            return

        wb = load_workbook(template_path, keep_vba=True)
        ws = wb['Payroll Reports Input']

        # Start writing at row 6 (row 5 has headers)
        start_row = 6

        for i, (idx, row) in enumerate(self.tagged_data.iterrows()):
            excel_row = start_row + i

            # Column E: Tag (Paycode allocation)
            ws.cell(row=excel_row, column=5, value=row.get('Tag'))

            # Column F: Code (Earning code)
            ws.cell(row=excel_row, column=6, value=row.get('Code'))

            # Column G: Hours
            ws.cell(row=excel_row, column=7, value=row.get('Hours'))

            # Column I: Amount
            ws.cell(row=excel_row, column=9, value=row.get('Amount'))

            # Column L: Taxes (tax code)
            if pd.notna(row.get('Tax_Code')):
                ws.cell(row=excel_row, column=12, value=row.get('Tax_Code'))

            # Column M: Amount (tax amount)
            if pd.notna(row.get('Tax_Amount')):
                ws.cell(row=excel_row, column=13, value=row.get('Tax_Amount'))

        wb.save(output_path)

    def get_questions(self) -> List[str]:
        """Return list of questions for user clarification"""
        # Remove duplicates
        return list(set(self.questions))
